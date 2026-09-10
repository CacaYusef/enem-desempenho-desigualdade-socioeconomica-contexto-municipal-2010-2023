"""Lossless column/row ingestion of ENEM; no sampling or eligibility filters."""

from __future__ import annotations

import argparse
import io
import json
import re
import time
import zipfile
from contextlib import contextmanager
from pathlib import Path

import duckdb
import pyarrow as pa
import pyarrow.csv as pacsv
import pyarrow.parquet as pq
from openpyxl import load_workbook

from enem_analysis.data.acquire import MANIFESTS, RAW, ROOT, sha256, write_json

PROCESSED = ROOT / "data" / "processed"
AUDIT = ROOT / "docs" / "sources" / "enem"


def source_path(year: int) -> Path:
    if year == 2023:
        return RAW / "microdados_enem_2023"
    return RAW / "enem" / str(year) / f"microdados_enem_{year}.zip"


@contextmanager
def package(year: int):
    path = source_path(year)
    if path.is_dir():
        files = {
            str(p.relative_to(path)).replace("\\", "/"): p
            for p in path.rglob("*")
            if p.is_file()
        }
        yield list(files), lambda name: files[name].open("rb")
    else:
        with zipfile.ZipFile(path) as archive:
            yield archive.namelist(), archive.open


def dictionary_sheets(content: bytes) -> dict:
    """Keep every official sheet separately, including ancillary questionnaires."""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    result = {}
    for sheet in workbook:
        variables = {}
        current = None
        name_index = 0
        raw_rows = []
        for row in sheet.iter_rows(values_only=True):
            values = [
                v if isinstance(v, (str, int, float, bool)) or v is None else str(v)
                for v in row
            ]
            raw_rows.append(values)
            for index, value in enumerate(values[:3]):
                if (
                    isinstance(value, str)
                    and value.strip().upper() == "NOME DA VARIÁVEL"
                ):
                    name_index = index
            if name_index >= len(values):
                continue
            name = values[name_index]
            if isinstance(name, str) and re.fullmatch(r"[A-Z][A-Z0-9_]*", name.strip()):
                if "_" in name or re.fullmatch(r"Q\d+", name):
                    current = name.strip()
                    variables[current] = {
                        "description": values[name_index + 1],
                        "categories": [],
                    }
            elif name is not None:
                current = None
            if current and len(values) > name_index + 3:
                code, label = values[name_index + 2 : name_index + 4]
                if code is not None:
                    variables[current]["categories"].append(
                        {"code": code, "label": label}
                    )
        result[sheet.title] = {"variables": variables, "rows": raw_rows}
    workbook.close()
    return result


def classify_geography(name: str, description: str) -> str:
    """Only declare residence when both the field and its dictionary support it."""
    label = description.lower()
    if (
        "resid" in label
        and "municip" in name.lower()
        and not name.endswith(("_ESC", "_PROVA"))
    ):
        return "residencia"
    if name.endswith("_PROVA") and ("prova" in label or "aplica" in label):
        return "prova"
    if name.endswith("_ESC") and "escola" in label:
        return "escola"
    return "nao_confirmado"


def import_year(year: int) -> dict:
    output = PROCESSED / "enem" / f"year={year}" / "participantes.parquet"
    audit_path = AUDIT / f"{year}.json"
    if output.exists() and audit_path.exists():
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        if audit.get("parquet_sha256") == sha256(output):
            print(f"EXISTS ENEM {year}: {audit['rows']} records", flush=True)
            return audit
        raise ValueError(f"Processed checksum mismatch: {output}")
    print(f"IMPORT ENEM {year}", flush=True)
    with package(year) as (names, opener):
        candidates = [
            n for n in names if n.upper().endswith(f"MICRODADOS_ENEM_{year}.CSV")
        ]
        if len(candidates) != 1:
            raise ValueError(f"Ambiguous main data for {year}: {candidates}")
        main = candidates[0]
        dictionaries = {}
        for name in names:
            if (
                name.lower().endswith(".xlsx")
                and "dicion" in name.lower()
                and not Path(name).name.startswith("~$")
            ):
                with opener(name) as stream:
                    dictionaries[name] = dictionary_sheets(stream.read())
        if not dictionaries:
            raise ValueError(f"No readable official XLSX dictionary: {year}")
        write_json(AUDIT / f"dictionary_{year}.json", dictionaries)
        with opener(main) as stream:
            header = stream.readline().decode("cp1252").strip().split(";")
        sheets = [
            sheet
            for dictionary in dictionaries.values()
            for sheet in dictionary.values()
        ]
        main_dictionary = max(
            sheets, key=lambda s: len(set(s["variables"]) & set(header))
        )
        meanings = main_dictionary["variables"]
        fields = []
        for name in header:
            if name.startswith("CO_MUNICIPIO"):
                description = str(meanings.get(name, {}).get("description") or "")
                fields.append(
                    {
                        "name": name,
                        "description": description,
                        "geography": classify_geography(name, description),
                    }
                )
        audit = {
            "year": year,
            "main_member": main,
            "files": names,
            "encoding": "cp1252",
            "delimiter": ";",
            "columns": header,
            "storage_types": "string; original empty fields represented as null",
            "filters": [],
            "rows_excluded": 0,
            "sampling": None,
            "sample_size_decided": False,
            "analysis_parameters": None,
            "municipality_fields": fields,
            "residence_available": any(f["geography"] == "residencia" for f in fields),
            "score_fields": [c for c in header if "NOTA" in c],
            "socioeconomic_fields": [c for c in header if re.fullmatch(r"Q\d+", c)],
            "variables_without_dictionary": sorted(set(header) - set(meanings)),
        }
        nulls = dict.fromkeys(header, 0)
        rows = 0
        output.parent.mkdir(parents=True, exist_ok=True)
        partial = output.with_suffix(".parquet.tmp")
        with opener(main) as stream:
            reader = pacsv.open_csv(
                stream,
                read_options=pacsv.ReadOptions(
                    encoding="cp1252", block_size=16 * 1024 * 1024
                ),
                parse_options=pacsv.ParseOptions(delimiter=";"),
                convert_options=pacsv.ConvertOptions(
                    column_types={c: pa.string() for c in header},
                    null_values=[""],
                    strings_can_be_null=True,
                ),
            )
            with pq.ParquetWriter(partial, reader.schema, compression="zstd") as writer:
                for batch in reader:
                    writer.write_batch(batch)
                    rows += batch.num_rows
                    for name, col in zip(header, batch.columns, strict=True):
                        nulls[name] += col.null_count
            if pq.ParquetFile(partial).metadata.num_rows != rows:
                raise ValueError("Row conservation failed")
        partial.replace(output)
    audit["rows"] = rows
    audit["missing_by_column"] = nulls
    connection = duckdb.connect()
    connection.execute("SET memory_limit='2GB'")
    relation = connection.read_parquet(str(output))
    relation.create_view("participants")
    if "NU_INSCRICAO" in header:
        nonnull, distinct = connection.execute(
            "SELECT count(NU_INSCRICAO), count(DISTINCT NU_INSCRICAO) FROM participants"
        ).fetchone()
        audit["duplicate_registration_excess"] = nonnull - distinct
    audit["observed_years"] = connection.execute(
        'SELECT "NU_ANO", count(*) FROM participants GROUP BY 1 ORDER BY 1'
    ).fetchall()
    audit["scores"] = {}
    for score in audit["score_fields"]:
        q = f'try_cast("{score}" AS DOUBLE)'
        valid, invalid, minimum, maximum, negatives = connection.execute(
            f'SELECT count({q}), count(*) FILTER (WHERE "{score}" IS NOT NULL '
            f"AND {q} IS NULL), "
            f"min({q}), max({q}), count(*) FILTER (WHERE {q}<0) FROM participants"
        ).fetchone()
        audit["scores"][score] = {
            "nonmissing_numeric": valid,
            "invalid_numeric": invalid,
            "min": minimum,
            "max": maximum,
            "negative": negatives,
        }
    municipalities = []
    for field in fields:
        code = field["name"]
        name = code.replace("CO_", "NO_", 1)
        name_sql = f'"{name}"' if name in header else "NULL::VARCHAR"
        distinct_codes = connection.execute(
            f'SELECT "{code}" codigo_original_ENEM, {name_sql} nome_original_ENEM, '
            "count(*) registros FROM participants GROUP BY 1,2"
        ).fetchdf()
        distinct_codes["ano_enem"] = year
        distinct_codes["campo_geografico"] = code
        distinct_codes["tipo_localizacao"] = field["geography"]
        municipalities.append(distinct_codes)
    if municipalities:
        import pandas as pd

        frame = pd.concat(municipalities, ignore_index=True)
        cross = PROCESSED / "enem" / f"year={year}" / "codigos_municipais.parquet"
        frame.to_parquet(cross, index=False, compression="zstd")
    connection.close()
    audit["parquet_bytes"] = output.stat().st_size
    audit["parquet_sha256"] = sha256(output)
    manifest_path = MANIFESTS / f"enem_{year}.json"
    audit["source_sha256"] = (
        json.loads(manifest_path.read_text(encoding="utf-8"))["sha256"]
        if manifest_path.exists()
        else "preexisting_2023_see_data_raw_README"
    )
    write_json(audit_path, audit)
    print(
        f"OK ENEM {year}: {rows:,} records, {len(header)} columns; "
        f"residence={audit['residence_available']}",
        flush=True,
    )
    return audit


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--years", type=int, nargs="+", default=list(range(2023, 2009, -1))
    )
    parser.add_argument("--wait-for-downloads", action="store_true")
    args = parser.parse_args()
    pending = set(args.years)
    errors = []
    deadline = time.monotonic() + 7200
    while pending:
        available = sorted(
            (y for y in pending if source_path(y).exists()), reverse=True
        )
        for year in available:
            try:
                import_year(year)
            except Exception as exc:
                errors.append({"year": year, "error": str(exc)})
                print(f"FAILED ENEM {year}: {exc}", flush=True)
            pending.remove(year)
        if not pending or not args.wait_for_downloads or time.monotonic() >= deadline:
            break
        status = PROCESSED / "acquisition_enem_status.json"
        if status.exists():
            failed_years = {
                int(e["id"].split("_")[-1])
                for e in json.loads(status.read_text())["errors"]
            }
            if pending.issubset(failed_years):
                break
        time.sleep(10)
    write_json(
        PROCESSED / "enem_import_status.json",
        {"errors": errors, "pending_years": sorted(pending)},
    )
    if errors or pending:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
